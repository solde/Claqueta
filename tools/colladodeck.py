"""Build a PowerPoint from a JSON spec with python-pptx, using the brand kit (colours, fonts). Usage: colladodeck.py spec.json out.pptx brand.json
Slide types: title, section, bullets, two, table, image, chart (bar|line|pie), quote. Every slide may carry "notes". Nested bullets: ["text", ["sub1", "sub2"]]."""
import json, sys, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

spec = json.load(open(sys.argv[1])); out = sys.argv[2]; brand = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else {}
def rgb(h, d="1f3a5f"): h = (h or d).lstrip("#"); h = (h * 2 if len(h) == 3 else h).ljust(6, "0")[:6]; return RGBColor.from_string(h.upper())
PRIM, ACC, INK, MUTED = rgb(brand.get("primary")), rgb(brand.get("accent"), "e0a100"), rgb(brand.get("text_color"), "1c1f24"), rgb(brand.get("muted"), "5b6472")
FH, FB = brand.get("font_heading") or "Calibri", brand.get("font_body") or "Calibri"
W, H = Inches(13.333), Inches(7.5)
prs = Presentation(); prs.slide_width, prs.slide_height = W, H
blank = prs.slide_layouts[6]
def text(slide, x, y, w, h, s, size=18, bold=False, color=INK, font=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = s; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font or FB; return tb
def rect(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(1, x, y, w, h); s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background(); return s
def chrome(slide, title, n):
    rect(slide, 0, 0, W, Inches(0.18), PRIM); rect(slide, 0, Inches(0.18), Inches(1.2), Inches(0.06), ACC)
    text(slide, Inches(0.6), Inches(0.45), Inches(12.1), Inches(0.9), title or "", 30, True, PRIM, FH, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, Inches(0.6), Inches(7.0), Inches(8), Inches(0.3), spec.get("footer", spec.get("title", "")), 10, False, MUTED)
    text(slide, Inches(11.7), Inches(7.0), Inches(1.1), Inches(0.3), str(n), 10, False, MUTED, align=PP_ALIGN.RIGHT)
def bullets(slide, x, y, w, h, items, size=20):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True; first = True
    def add(it, lvl):
        nonlocal first
        if isinstance(it, list):
            for s in it: add(s, lvl + 1)
            return
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False; p.level = lvl; p.space_after = Pt(8)
        r = p.add_run(); r.text = ("• " if lvl == 0 else "– ") + str(it); r.font.size = Pt(size - 3 * lvl); r.font.color.rgb = INK; r.font.name = FB
    for it in items: add(it, 0)
for i, s in enumerate(spec.get("slides", []), start=1):
    sl = prs.slides.add_slide(blank); t = s.get("type", "bullets")
    if t == "title":
        rect(sl, 0, 0, W, H, PRIM); rect(sl, Inches(0.8), Inches(3.55), Inches(1.6), Inches(0.09), ACC)
        text(sl, Inches(0.8), Inches(1.9), Inches(11.5), Inches(1.6), s.get("title", spec.get("title", "")), 44, True, RGBColor(255, 255, 255), FH, anchor=MSO_ANCHOR.BOTTOM)
        text(sl, Inches(0.8), Inches(3.9), Inches(11.5), Inches(1.0), s.get("subtitle", spec.get("subtitle", "")), 22, False, RGBColor(235, 235, 235))
        text(sl, Inches(0.8), Inches(6.4), Inches(11.5), Inches(0.5), s.get("meta", spec.get("meta", "")), 14, False, RGBColor(220, 220, 220))
    elif t == "section":
        rect(sl, 0, 0, W, H, ACC); text(sl, Inches(0.8), Inches(2.8), Inches(11.5), Inches(1.6), s.get("title", ""), 40, True, INK, FH, anchor=MSO_ANCHOR.MIDDLE)
    elif t == "quote":
        chrome(sl, s.get("title", ""), i); text(sl, Inches(1.2), Inches(2.2), Inches(10.9), Inches(3), "“" + s.get("text", "") + "”", 30, False, PRIM, FH, anchor=MSO_ANCHOR.MIDDLE)
        if s.get("by"): text(sl, Inches(1.2), Inches(5.4), Inches(10.9), Inches(0.6), "— " + s["by"], 16, False, MUTED)
    elif t == "two":
        chrome(sl, s.get("title", ""), i); bullets(sl, Inches(0.6), Inches(1.6), Inches(5.9), Inches(5.2), s.get("left", []), 18); bullets(sl, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.2), s.get("right", []), 18)
    elif t == "table":
        chrome(sl, s.get("title", ""), i); cols, rows = s.get("columns", []), s.get("rows", []); n = max(1, len(rows)) + 1
        gt = sl.shapes.add_table(n, max(1, len(cols)), Inches(0.6), Inches(1.6), Inches(12.1), Inches(min(5.2, 0.45 * n))).table
        for c, name in enumerate(cols):
            cell = gt.cell(0, c); cell.text = str(name); cell.fill.solid(); cell.fill.fore_color.rgb = PRIM
            for p in cell.text_frame.paragraphs:
                for r in p.runs: r.font.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(255, 255, 255); r.font.name = FH
        for ri, row in enumerate(rows, start=1):
            for c, v in enumerate(row):
                cell = gt.cell(ri, c); cell.text = str(v); cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(244, 246, 248) if ri % 2 == 0 else RGBColor(255, 255, 255)
                for p in cell.text_frame.paragraphs:
                    for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = INK; r.font.name = FB
    elif t == "image":
        chrome(sl, s.get("title", ""), i); p = s.get("path", "")
        if os.path.exists(p): sl.shapes.add_picture(p, Inches(0.9), Inches(1.5), height=Inches(5.0) if not s.get("full") else None, width=Inches(11.5) if s.get("full") else None)
        if s.get("caption"): text(sl, Inches(0.9), Inches(6.55), Inches(11.5), Inches(0.4), s["caption"], 12, False, MUTED, align=PP_ALIGN.CENTER)
    elif t == "chart":
        chrome(sl, s.get("title", ""), i); cd = CategoryChartData(); cd.categories = s.get("categories", [])
        for se in s.get("series", []): cd.add_series(se.get("name", ""), se.get("values", []))
        kind = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE, "hbar": XL_CHART_TYPE.BAR_CLUSTERED}.get(s.get("chart", "bar"), XL_CHART_TYPE.COLUMN_CLUSTERED)
        ch = sl.shapes.add_chart(kind, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3), cd).chart; ch.has_legend = len(s.get("series", [])) > 1 or kind == XL_CHART_TYPE.PIE
        if ch.has_legend: ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        pal = [PRIM, ACC, rgb("2f9e73"), rgb("d1495b"), rgb("6c5ce7"), rgb("00a6c0")]
        if kind != XL_CHART_TYPE.PIE:
            for j, ser in enumerate(ch.series): ser.format.fill.solid(); ser.format.fill.fore_color.rgb = pal[j % len(pal)]
    else:
        chrome(sl, s.get("title", ""), i); bullets(sl, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2), s.get("bullets", []), 22)
    if s.get("notes"): sl.notes_slide.notes_text_frame.text = s["notes"]
prs.core_properties.title = spec.get("title", ""); prs.core_properties.author = spec.get("author", "Collado"); prs.save(out); print(out)
