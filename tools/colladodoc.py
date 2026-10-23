"""Helpers for building Word/Excel files by hand with python-docx / openpyxl in the team's brand (colours, fonts, margins) when the Markdown route is not enough.
    import sys; sys.path.insert(0, "/opt/collado-tools"); from colladodoc import brand, new_docx, add_table, new_xlsx, write_table
    b = brand()                       # default brand from the app (or brand("Sobrio"))
    doc = new_docx(b); doc.add_heading("Título", 1); add_table(doc, ["A", "B"], [[1, 2]], b); doc.save("x.docx")
    wb, ws = new_xlsx(b); write_table(ws, ["Concepto", "Importe"], [["Cámara", 1200.5]], b, money=[1]); wb.save("x.xlsx")"""
import json, os, urllib.request
def _api(path):
    tok = open(os.environ.get("COLLADO_TOKEN_FILE") or os.path.expanduser("~/.collado/token")).read().strip()
    req = urllib.request.Request(os.environ.get("COLLADO_API", "http://127.0.0.1:3000") + path, headers={"authorization": "Bearer " + tok, "x-requested-with": "collado"})
    with urllib.request.urlopen(req, timeout=30) as r: return json.load(r)
def brand(name=None):
    rows = _api("/api/e/brands?limit=100")["rows"]
    b = next((x for x in rows if x["name"] == name), None) if name else next((x for x in rows if x.get("is_default")), rows[0] if rows else {})
    return b or {}
def _rgb(h, d="1f3a5f"):
    from docx.shared import RGBColor
    h = (h or d).lstrip("#"); return RGBColor.from_string((h * 2 if len(h) == 3 else h).ljust(6, "0")[:6].upper())
def new_docx(b):
    from docx import Document
    from docx.shared import Pt, Mm
    d = Document(); s = d.sections[0]; m = (b.get("margin_mm") or {}); s.top_margin, s.bottom_margin = Mm(m.get("top", 25)), Mm(m.get("bottom", 22)); s.left_margin, s.right_margin = Mm(m.get("left", 20)), Mm(m.get("right", 20))
    n = d.styles["Normal"]; n.font.name = b.get("font_body") or "Calibri"; n.font.size = Pt(b.get("base_size") or 10.5)
    for lvl, size in ((1, 24), (2, 16), (3, 13)):
        st = d.styles[f"Heading {lvl}"]; st.font.name = b.get("font_heading") or "Calibri"; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = _rgb(b.get("primary"))
    return d
def add_table(doc, columns, rows, b, widths=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    t = doc.add_table(rows=1, cols=len(columns)); t.style = "Table Grid"
    for i, c in enumerate(columns):
        cell = t.rows[0].cells[i]; cell.text = str(c); tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), (b.get("table_head_bg") or b.get("primary") or "#1f3a5f").lstrip("#")); tcPr.append(shd)
        for r in cell.paragraphs[0].runs: r.font.bold = True; r.font.color.rgb = _rgb("ffffff")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = "" if v is None else str(v)
    return t
def new_xlsx(b):
    import openpyxl
    wb = openpyxl.Workbook(); return wb, wb.active
def write_table(ws, columns, rows, b, money=(), pct=(), date=(), total=()):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    head = PatternFill("solid", fgColor=(b.get("table_head_bg") or b.get("primary") or "#1f3a5f").lstrip("#"))
    ws.append(list(columns))
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF", name=b.get("font_body") or "Calibri"); c.fill = head; c.alignment = Alignment(vertical="center")
    for r in rows: ws.append(list(r))
    last = ws.max_row
    for i in money: 
        for c in ws[get_column_letter(i + 1)][1:]: c.number_format = '#,##0.00 "€"'
    for i in pct:
        for c in ws[get_column_letter(i + 1)][1:]: c.number_format = "0.0%"
    for i in date:
        for c in ws[get_column_letter(i + 1)][1:]: c.number_format = "dd/mm/yyyy"
    if total:
        ws.append(["Total"] + [None] * (len(columns) - 1)); tr = ws.max_row
        for i in total: ws.cell(tr, i + 1, f"=SUM({get_column_letter(i + 1)}2:{get_column_letter(i + 1)}{last})").font = Font(bold=True)
    for i, c in enumerate(columns, start=1): ws.column_dimensions[get_column_letter(i)].width = min(60, max(12, len(str(c)) + 4, *[len(str(r[i - 1])) + 2 for r in rows if i - 1 < len(r) and r[i - 1] is not None] or [0]))
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
